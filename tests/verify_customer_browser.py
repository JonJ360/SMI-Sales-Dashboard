"""Frozen local snapshot/browser/XLSX validation. Never signs in or publishes data."""
import argparse
import json
from pathlib import Path
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import threading
from decimal import Decimal
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]

def run(payload,output):
    output.mkdir(parents=True,exist_ok=True)
    raw=payload.read_bytes();data=json.loads(raw);checks=[];errors=[]
    class Handler(SimpleHTTPRequestHandler):
        def log_message(self,*args): pass
        def do_GET(self):
            if self.path.split('?')[0]=='/data/sales.json':
                self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(raw)
            else: super().do_GET()
    server=ThreadingHTTPServer(('127.0.0.1',0),partial(Handler,directory=str(ROOT)))
    worker=threading.Thread(target=server.serve_forever,daemon=True);worker.start()
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(headless=True,channel='msedge')
            page=browser.new_page(viewport={'width':1440,'height':1000},accept_downloads=True)
            page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto(f'http://127.0.0.1:{server.server_port}',wait_until='networkidle');page.wait_for_function('!!state.data')
            page.locator('[data-view=customers]').click()
            assert page.locator('#customersPeriod').input_value()=='YTD'
            visible=page.locator('#customersBody tr td:nth-child(2)').all_text_contents()
            assert visible==[x['name'] for x in data['rankings']['YTD']['customers']], 'Default must be actual current top25, not prior cohort'
            outside=set(visible)-{x['name'] for x in data['comparisons']['YTD']['customer_comparison']}
            assert outside,'Real snapshot must exercise current customers outside prior25'
            for period,month in [('YTD',''),('1M',''),('FULL',''),('MONTH','2024-01'),('MONTH',data['as_of'][:7])]:
                page.locator('#customersPeriod').select_option(period)
                if month: page.locator('#customersMonth').select_option(month)
                for cohort in ['current','prior']:
                    page.locator('#customerCohort').select_option(cohort)
                    selected=data['months'][month] if period=='MONTH' else data['comparisons'][period]
                    ranking=selected['rankings'] if period=='MONTH' else data['rankings'][period]
                    source=selected['customer_comparison'] if cohort=='prior' else ranking['customers']
                    assert page.locator('#customersBody tr td:nth-child(2)').all_text_contents()==[x['name'] for x in source]
                    target=next((x for x in source if x['name'] in outside),source[0])
                    page.locator('#customerSearch').fill(target['name'].lower())
                    expected=[(i,x) for i,x in enumerate(source,1) if target['name'].lower() in x['name'].lower()]
                    assert page.locator('#customersBody tr.clickable').count()==len(expected)
                    assert page.locator('#customersBody tr td').first.inner_text()==str(expected[0][0])
                    page.locator('#customersBody tr.clickable').first.click()
                    assert page.locator('#customerDrawerName').inner_text()==expected[0][1]['name']
                    assert 'No comparable value' not in page.locator('#customerDrawerPeriod').inner_text()
                    page.locator('#customerDrawerClose').click()
                    # Decode an actual browser-downloaded workbook independently of SheetJS.
                    with page.expect_download() as dl: page.locator('#customerExport').click()
                    path=output/f'{period}-{month or "all"}-{cohort}.xlsx';dl.value.save_as(path)
                    wb=load_workbook(path,data_only=False);assert wb.sheetnames==['Customer Comparison','Report Info']
                    rows=list(wb['Customer Comparison'].values)[1:];assert len(rows)==len(expected)
                    for actual,(rank,x) in zip(rows,expected):
                        current=x['current_sales'] if cohort=='prior' else x['sales'];prior=x['prior_sales']
                        delta=Decimal(str(current))-Decimal(str(prior))
                        assert actual[:4]==(rank,x['name'],prior,current)
                        assert Decimal(str(actual[4]))==delta
                        if prior==0: assert actual[5]=='N/A'
                        else: assert abs(actual[5]-float(delta)/abs(prior))<1e-12
                    info=dict(list(wb['Report Info'].values)[1:]);assert info['Source snapshot SHA-256']==data['sha256'];assert info['Exported row count']==len(rows)
                    assert info['Data as of'].date().isoformat()==data['as_of'];assert info['Refresh / publication heartbeat (UTC)']==data['refreshed_at']
                    assert info['Period']==period;assert info['Ranking cohort']==('Prior-period Top 25' if cohort=='prior' else 'Current-period Top 25')
                    checks.append({'period':period,'month':month,'cohort':cohort,'rows':len(rows),'workbook':path.name})
                    page.locator('#customerSearch').fill('')
            page.locator('#customersPeriod').select_option('YTD');page.locator('#customerCohort').select_option('current')
            page.screenshot(path=str(output/'customers-desktop.png'),full_page=True)
            page.set_viewport_size({'width':390,'height':844});page.screenshot(path=str(output/'customers-mobile.png'),full_page=True)
            assert page.locator('#customers .table-tools').evaluate('(e)=>e.scrollWidth<=e.clientWidth'), 'Customer controls must fit mobile panel'
            assert page.locator('#customerSearch').evaluate('(e)=>e.getBoundingClientRect().width>=150'), 'Search must remain usable'
            assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
            page.set_viewport_size({'width':1440,'height':1000})
            # Export must use the rendered immutable report, not mutable refreshed state.
            before=page.evaluate('state.customerReport.source_sha256')
            page.evaluate("()=>{state.data={...state.data,sha256:'changed-after-render'};}")
            with page.expect_download() as dl: page.locator('#customerExport').click()
            path=output/'frozen-visible-export.xlsx';dl.value.save_as(path)
            assert dict(list(load_workbook(path)['Report Info'].values)[1:])['Source snapshot SHA-256']==before
            # Literal punctuation and formula prefixes remain text and do not execute HTML/JS.
            hostile="=O'Brien <img src=x onerror=alert(1)>"
            page.evaluate('(name)=>{state.data.rankings.YTD.customers=[{name,sales:-50.25,prior_sales:-100.5},{name:"Zero",sales:1,prior_sales:0}];renderCustomers()}',hostile)
            assert page.locator('#customersBody img').count()==0
            assert '50.0%' in page.locator('#customersBody').inner_text();assert 'N/A' in page.locator('#customersBody').inner_text()
            page.locator('#customersBody tr.clickable').first.click() # unknown detail is safely ignored
            with page.expect_download() as dl: page.locator('#customerExport').click()
            path=output/'safe-text.xlsx';dl.value.save_as(path);cell=load_workbook(path)['Customer Comparison']['B2'];assert cell.data_type=='s' and cell.value==hostile
            page.locator('#customerSearch').fill('no match xyz');assert 'No customers' in page.locator('#customersBody').inner_text()
            with page.expect_download() as dl: page.locator('#customerExport').click()
            path=output/'empty.xlsx';dl.value.save_as(path);assert load_workbook(path)['Customer Comparison'].max_row==1
            assert not errors,errors
            browser.close()
    finally: server.shutdown();server.server_close();worker.join()
    result={'mode':'local preview with frozen real payload; not authenticated live E2E','checks':checks,'outside_prior_cohort_count':len(outside),'page_errors':errors,'frozen_export':True,'safe_text':True,'empty_export':True}
    (output/'browser-validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8');print(json.dumps(result))
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--payload',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);a=ap.parse_args();run(a.payload,a.output)
